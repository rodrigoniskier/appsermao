import re
from io import BytesIO

from docx import Document
from docx.shared import Pt
from flask import Blueprint, abort, redirect, render_template, send_file, url_for
from flask_login import current_user, login_required

from exposibot.extensions import db
from exposibot.models import Sermon

bp = Blueprint("dashboard", __name__)


def _get_owned_sermon(sermon_id):
    sermon = db.session.get(Sermon, sermon_id)
    if not sermon or sermon.user_id != current_user.id:
        abort(404)
    return sermon


def _render_workspace(sermon):
    """Render the workspace and add the research export control without duplicating template chrome."""
    html = render_template("workspace.html", sermon=sermon)
    marker = '<div class="action-area">'
    export_url = url_for("dashboard.download_research_docx", sermon_id=sermon.id)
    export_control = (
        f'{marker}<div class="form-group">'
        f'<a class="btn btn-muted full-width" href="{export_url}" '
        'target="_blank" rel="noopener">Exportar pesquisa (.docx)</a>'
        "</div>"
    )
    return html.replace(marker, export_control, 1)


@bp.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard.dashboard"))
    return redirect(url_for("auth.login"))


@bp.route("/dashboard")
@login_required
def dashboard():
    sermons = (
        Sermon.query.filter_by(user_id=current_user.id)
        .order_by(Sermon.updated_at.desc())
        .all()
    )
    return render_template("dashboard.html", sermons=sermons)


@bp.route("/sermon/new", methods=["POST"])
@login_required
def new_sermon():
    sermon = Sermon(user_id=current_user.id, research_notes=[], outline={})
    db.session.add(sermon)
    db.session.commit()
    return redirect(url_for("dashboard.workspace", sermon_id=sermon.id))


@bp.route("/sermon/<int:sermon_id>")
@login_required
def workspace(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    return _render_workspace(sermon)


@bp.route("/sermon/<int:sermon_id>/delete", methods=["POST"])
@login_required
def delete_sermon(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    db.session.delete(sermon)
    db.session.commit()
    return redirect(url_for("dashboard.dashboard"))


@bp.route("/sermon/<int:sermon_id>/preview")
@login_required
def preview(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    outline = sermon.outline or {}
    return render_template(
        "sermon_preview.html",
        sermon=sermon,
        outline=outline,
        topicos=outline.get("topicos", []),
    )


def _build_docx(sermon):
    outline = sermon.outline or {}
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    document.add_heading("Esboço de Pregação Expositiva", 0)
    document.add_paragraph(f"Texto Base: {sermon.reference}")
    document.add_paragraph(f"ICT: {outline.get('ict', '')}")
    document.add_paragraph(f"Tese: {outline.get('tese', '')}")
    if outline.get("fcd"):
        document.add_paragraph(f"FCD: {outline.get('fcd', '')}")
    document.add_paragraph(
        f"Propósito: {outline.get('proposito_basico', '')} | {outline.get('proposito_especifico', '')}"
    )
    if outline.get("proposito_redentivo"):
        document.add_paragraph(f"Propósito redentivo: {outline.get('proposito_redentivo', '')}")
    document.add_paragraph("-" * 60)

    document.add_heading("Introdução", level=1)
    document.add_paragraph(outline.get("intro", ""))

    def _bold_label(text):
        document.add_paragraph().add_run(text).bold = True

    for i, topico in enumerate(outline.get("topicos", []), start=1):
        document.add_heading(f"{i}. {topico.get('titulo', '')}", level=2)
        if topico.get("texto_base"):
            document.add_paragraph(f"Texto: {topico.get('texto_base', '')}")
        _bold_label("Elucidação:")
        document.add_paragraph(topico.get("explicacao", ""))
        if topico.get("ilustracao"):
            _bold_label("Ilustração:")
            p = document.add_paragraph(topico["ilustracao"])
            try:
                p.style = document.styles["Quote"]
            except KeyError:
                pass
        _bold_label("Aplicação:")
        document.add_paragraph(topico.get("aplicacao", ""))
        if topico.get("transicao"):
            _bold_label("Transição:")
            document.add_paragraph(topico.get("transicao", ""))

    if outline.get("conexao_cristocentrica"):
        document.add_heading("Conexão Cristocêntrica", level=1)
        document.add_paragraph(outline.get("conexao_cristocentrica", ""))

    document.add_heading("Conclusão", level=1)
    document.add_paragraph(outline.get("conclusao", ""))

    buf = BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


def _add_inline_markdown(paragraph, text):
    """Render the small inline Markdown subset commonly returned by research providers."""
    token_pattern = re.compile(
        r"(\*\*.+?\*\*|__.+?__|`.+?`|\*[^*]+?\*|_[^_]+?_|\[[^\]]+\]\([^)]+\))"
    )
    cursor = 0
    for match in token_pattern.finditer(text or ""):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])

        token = match.group(0)
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
        if link:
            paragraph.add_run(link.group(1))
            paragraph.add_run(f" ({link.group(2)})")
        elif token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()

    if cursor < len(text or ""):
        paragraph.add_run(text[cursor:])


def _add_markdown_to_docx(document, markdown_text):
    """Convert research Markdown to readable Word paragraphs without injecting HTML."""
    pending = []

    def flush_pending():
        if not pending:
            return
        text = " ".join(part.strip() for part in pending if part.strip())
        pending.clear()
        if text:
            paragraph = document.add_paragraph()
            _add_inline_markdown(paragraph, text)

    normalized = str(markdown_text or "").replace("\r\n", "\n").replace("\r", "\n")
    for raw_line in normalized.split("\n"):
        line = raw_line.strip()
        if not line:
            flush_pending()
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        ordered = re.match(r"^\d+[.)]\s+(.+)$", line)

        if heading:
            flush_pending()
            level = min(len(heading.group(1)) + 1, 3)
            paragraph = document.add_heading(level=level)
            _add_inline_markdown(paragraph, heading.group(2))
        elif bullet:
            flush_pending()
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_markdown(paragraph, bullet.group(1))
        elif ordered:
            flush_pending()
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_markdown(paragraph, ordered.group(1))
        elif line.startswith(">"):
            flush_pending()
            paragraph = document.add_paragraph(style="Quote")
            _add_inline_markdown(paragraph, line.lstrip("> "))
        elif re.fullmatch(r"[-*_]{3,}", line):
            flush_pending()
        else:
            pending.append(line)

    flush_pending()


def _build_research_docx(sermon):
    notes = sermon.research_notes or []
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    document.add_heading("Material de Pesquisa", 0)

    sermon_line = document.add_paragraph()
    sermon_line.add_run("Sermão: ").bold = True
    sermon_line.add_run(sermon.display_title())

    reference_line = document.add_paragraph()
    reference_line.add_run("Texto base: ").bold = True
    reference_line.add_run(sermon.reference or "Não informado")

    count_line = document.add_paragraph()
    count_line.add_run("Itens coletados: ").bold = True
    count_line.add_run(str(len(notes)))

    if not notes:
        document.add_paragraph("Nenhum material de pesquisa foi incorporado até o momento.")
    else:
        for index, note in enumerate(notes, start=1):
            tipo = str(note.get("tipo") or "Nota de pesquisa")
            document.add_heading(f"{index}. {tipo}", level=1)
            _add_markdown_to_docx(document, note.get("markdown", ""))

    buf = BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


@bp.route("/sermon/<int:sermon_id>/download.docx")
@login_required
def download_docx(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    buf = _build_docx(sermon)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"sermao_{sermon.id}.docx",
    )


@bp.route("/sermon/<int:sermon_id>/research.docx")
@login_required
def download_research_docx(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    buf = _build_research_docx(sermon)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"pesquisa_sermao_{sermon.id}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@bp.route("/sermon/<int:sermon_id>/download.md")
@login_required
def download_md(sermon_id):
    sermon = _get_owned_sermon(sermon_id)
    outline = sermon.outline or {}

    lines = [
        f"# {sermon.display_title()}",
        "",
        f"**Texto Base:** {sermon.reference}",
        f"**ICT:** {outline.get('ict', '')}",
        f"**Tese:** {outline.get('tese', '')}",
    ]
    if outline.get("fcd"):
        lines.append(f"**FCD:** {outline.get('fcd', '')}")
    lines.append(
        f"**Propósito:** {outline.get('proposito_basico', '')} — {outline.get('proposito_especifico', '')}"
    )
    if outline.get("proposito_redentivo"):
        lines.append(f"**Propósito redentivo:** {outline.get('proposito_redentivo', '')}")

    lines += [
        "",
        "## Introdução",
        outline.get("intro", ""),
        "",
        "## Desenvolvimento",
    ]
    for i, topico in enumerate(outline.get("topicos", []), start=1):
        lines.append(f"### {i}. {topico.get('titulo', '')}")
        if topico.get("texto_base"):
            lines.append(f"**Texto:** {topico.get('texto_base', '')}")
        lines.append(f"**Elucidação:** {topico.get('explicacao', '')}")
        if topico.get("ilustracao"):
            lines.append(f"> 💡 Ilustração: {topico['ilustracao']}")
        lines.append(f"**Aplicação:** {topico.get('aplicacao', '')}")
        if topico.get("transicao"):
            lines.append(f"**Transição:** {topico.get('transicao', '')}")
        lines.append("")

    if outline.get("conexao_cristocentrica"):
        lines += [
            "## Conexão Cristocêntrica",
            outline.get("conexao_cristocentrica", ""),
            "",
        ]

    lines += ["## Conclusão", outline.get("conclusao", "")]

    buf = BytesIO("\n\n".join(lines).encode("utf-8"))
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"sermao_{sermon.id}.md",
        mimetype="text/markdown",
    )
