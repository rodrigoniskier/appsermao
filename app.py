import os
import json
import markdown
import time
import re
from flask import Flask, render_template, request, jsonify, send_file
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt
from io import BytesIO

# --- BIBLIOTECAS DE IA E BUSCA ---
from groq import Groq
from tavily import TavilyClient

# --- 1. CONFIGURAÇÃO E SEGURANÇA ---

# Define caminho absoluto para o arquivo .env (Essencial no PythonAnywhere)
project_folder = os.path.expanduser('~/mysite')
load_dotenv(os.path.join(project_folder, '.env'))

# Recupera as chaves de API
groq_api_key = os.getenv("GROQ_API_KEY")
tavily_api_key = os.getenv("TAVILY_API_KEY")

# Inicializa Clientes (com tratamento de erro básico na inicialização)
groq_client = None
tavily_client = None

if groq_api_key:
    groq_client = Groq(api_key=groq_api_key)
else:
    print("AVISO: GROQ_API_KEY não encontrada no .env")

if tavily_api_key:
    tavily_client = TavilyClient(api_key=tavily_api_key)
else:
    print("AVISO: TAVILY_API_KEY não encontrada no .env")

app = Flask(__name__)

# --- 2. SYSTEM PROMPTS (A PERSONALIDADE DO BOT) ---

MASTER_SYSTEM_PROMPT = """
Você é o 'Exposibot', um Motor de Pesquisa Teológica e Mentor de Pregação Expositiva.
Você opera sob um estrito **Protocolo de Integridade Informacional Absoluta (IIA)**.

### 1. FUNDAMENTO TEOLÓGICO (FILTRO EXCLUSIVO)
Sua cosmovisão e base de dados devem se restringir estritamente à **Teologia Reformada Calvinista e Neo-Calvinista**.
* **Fontes Aceitáveis:** Autores clássicos (Calvino, Turretin, Bavinck, Kuyper) e contemporâneos fiéis (Keller, Beale, Carson, Sproul, Piper, Lloyd-Jones).
* **Fontes a Ignorar:** Teologia Liberal, Arminiosismo, Teologia da Prosperidade ou exegese Católica Romana.
* **Padrão Confessional:** Padrões de Westminster.

### 2. PROTOCOLO DE INTEGRIDADE INFORMACIONAL (IIA)
* **Zero Alucinação:** Nunca invente dados. Se o contexto da pesquisa não fornecer a resposta, diga: "Não encontrei dados suficientes nas fontes reformadas consultadas".
* **Referenciamento Obrigatório:** Toda afirmação factual deve ser rastreável.
* **Citação de Fontes:** Você DEVE finalizar a resposta com uma seção chamada `### Referências Consultadas`, listando os Títulos e URLs exatos fornecidos no contexto da pesquisa.

### 3. ESTILO DE RESPOSTA
* Acadêmico, denso, porém pastoral.
* Estruturado em Markdown.
* Foco na centralidade de Cristo (Hermenêutica Redentiva-Histórica).
"""

HOMILETICS_SYSTEM_PROMPT = """
Você é um Professor de Homilética Reformada (linha Bryan Chapell / Haddon Robinson).
Sua tarefa é estruturar um esboço de sermão expositivo com base ESTRITAMENTE nas notas de pesquisa fornecidas e no texto bíblico.
SAÍDA OBRIGATÓRIA: Apenas um objeto JSON válido.
"""

# --- 3. FUNÇÕES AUXILIARES (BUSCA E GERAÇÃO) ---

def perform_grounded_search(query):
    """
    Usa o Tavily para buscar fontes confiáveis na web.
    Força a busca em domínios ou contextos reformados.
    """
    if not tavily_client:
        return "Aviso: Pesquisa na web desativada (Tavily Key ausente)."

    try:
        # Adiciona sufixos para garantir qualidade teológica na busca
        reformed_query = f"{query} (site:ligonier.org OR site:monergism.com OR site:thegospelcoalition.org OR site:9marks.org OR site:desiringgod.org OR site:puritanboard.com OR site:banneroftruth.org OR site:reformed.org OR \"Reformed Theology\" OR \"Calvinism\")"

        # Busca avançada
        response = tavily_client.search(
            query=reformed_query,
            search_depth="advanced",
            max_results=5,
            include_answer=False,
            include_raw_content=False
        )

        context_str = "### CONTEXTO DAS FONTES (Estritamente Reformadas):\n"
        for result in response.get('results', []):
            context_str += f"- Título: {result['title']}\n"
            context_str += f"  Link Real: {result['url']}\n"
            context_str += f"  Resumo: {result['content']}\n\n"

        return context_str
    except Exception as e:
        print(f"Erro no Tavily: {e}")
        return f"Erro ao realizar pesquisa externa: {str(e)}"

def generate_with_groq(prompt, context, system_prompt=MASTER_SYSTEM_PROMPT, json_mode=False):
    """
    Gera a resposta usando Groq (Llama 3).
    Suporta troca de System Prompt e JSON Mode.
    """
    if not groq_client:
        return "Erro: Cliente Groq não configurado."

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": f"CONTEXTO DE DADOS:\n{context}\n\n---\n\nCOMANDO:\n{prompt}"}
    ]

    kwargs = {
        "messages": messages,
        "model": "llama-3.3-70b-versatile",
        "temperature": 0.1, # Temperatura baixa para fidelidade
        "max_tokens": 4096,
    }

    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}

    try:
        chat_completion = groq_client.chat.completions.create(**kwargs)
        return chat_completion.choices[0].message.content
    except Exception as e:
        return f"Erro na geração Groq: {str(e)}"

# --- 4. ROTAS DE NAVEGAÇÃO (PÁGINAS) ---

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/homiletics', methods=['POST'])
def homiletics():
    notes = request.form.get('notes_content', '')
    reference = request.form.get('reference', '')
    return render_template('homiletics.html', notes=notes, reference=reference)

@app.route('/view_sermon', methods=['POST'])
def view_sermon():
    data = request.form
    topicos = []
    titulos = data.getlist('topico_titulo[]')
    explicacoes = data.getlist('topico_explicacao[]')
    ilustracoes = data.getlist('topico_ilustracao[]')
    aplicacoes = data.getlist('topico_aplicacao[]')

    for i in range(len(titulos)):
        topicos.append({
            'titulo': titulos[i],
            'explicacao': explicacoes[i],
            'ilustracao': ilustracoes[i],
            'aplicacao': aplicacoes[i]
        })
    return render_template('sermon.html', data=data, topicos=topicos)

# --- 5. ROTA DA API DE PESQUISA (ORQUESTRAÇÃO: GROQ + TAVILY) ---

@app.route('/api/analyze', methods=['POST'])
def analyze_text():
    if not groq_client:
        return jsonify({"error": "Erro no Servidor: API Key do Groq não configurada."}), 500

    try:
        data = request.json
        texto = data.get('texto') # Referência Bíblica
        tipo = data.get('tipo')   # Tipo de Análise

        if not texto:
            return jsonify({"error": "A referência bíblica está vazia."}), 400

        base_context = f"{texto} reformed theology calvinist commentary"

        # Mapeia botões para queries de busca otimizadas
        search_queries = {
            "Dados Geográficos": f"geography archaeology location {texto} biblical scholar commentary",
            "Perfil Biográfico": f"biography characters {texto} redemptive history reformed",
            "Contexto Histórico": f"historical context authorship date {texto} conservative commentary carson moo",
            "Contexto Canônico": f"biblical theology canon connections {texto} beale vos",
            "Estilo Literário": f"literary genre structure chiasm {texto} analysis",
            "Contexto Político": f"political context roman empire {texto} historical background",
            "Sócio-Cultural": f"social cultural customs {texto} ancient near east background",
            "Contexto Filosófico": f"philosophical background heresy {texto} biblical worldview",
            "Teologia Sistemática": f"systematic theology doctrines {texto} westminster confession berkhof",
            "Foco Cristocêntrico": f"christ centered exposition {texto} keller chapell",
            "Exegese (BHS/UBS5)": f"exegesis greek hebrew syntax {texto} technical commentary"
        }

        query = search_queries.get(tipo, base_context)
        search_context = perform_grounded_search(query)

        base_prompt = f"Com base na referência '{texto}' e ESTRITAMENTE no contexto fornecido:"

        # Prompts estruturados para o LLM
        prompts = {
            "Dados Geográficos": f"{base_prompt} Atue como Arqueólogo. Detalhe localização, topografia, clima e teologia do lugar.",
            "Perfil Biográfico": f"{base_prompt} Faça o perfil biográfico, etimologia e papel na redenção dos personagens.",
            "Contexto Histórico": f"{base_prompt} Detalhe autoria, data, destinatários e propósito (visão conservadora/reformada).",
            "Contexto Canônico": f"{base_prompt} Situe no cânon, conexões pactuais e referências cruzadas.",
            "Estilo Literário": f"{base_prompt} Analise gênero, estrutura (quiasmo?), figuras de linguagem e tom.",
            "Contexto Político": f"{base_prompt} Analise estruturas de poder e tensões políticas da época.",
            "Sócio-Cultural": f"{base_prompt} Explique costumes, leis sociais e impedimentos de leitura.",
            "Contexto Filosófico": f"{base_prompt} Identifique cosmovisões em choque e vocabulário filosófico.",
            "Teologia Sistemática": f"{base_prompt} Categorize sistematicamente e conecte com a Confissão de Fé de Westminster.",
            "Foco Cristocêntrico": f"{base_prompt} Aplique hermenêutica redentiva-histórica para pregar Cristo.",
            "Exegese (BHS/UBS5)": f"{base_prompt} Faça exegese técnica (original, morfologia, sintaxe). Use transliteração."
        }

        user_instruction = prompts.get(tipo, f"Analise teologicamente: {texto}")
        user_instruction += "\n\nIMPORTANTE: Finalize com a lista '### Referências Consultadas' contendo os links reais fornecidos no contexto."

        final_response = generate_with_groq(user_instruction, search_context)
        resposta_html = markdown.markdown(final_response)

        return jsonify({"resultado": resposta_html})

    except Exception as e:
        print(f"ERRO API: {str(e)}")
        return jsonify({"error": f"Erro interno: {str(e)}"}), 500


# --- 6. NOVA ROTA: SUGESTÃO DE SERMÃO (JSON) ---

@app.route('/api/suggest_sermon', methods=['POST'])
def suggest_sermon():
    if not groq_client: return jsonify({"error": "Groq off"}), 500
    try:
        data = request.json
        notes = data.get('notes', '')
        reference = data.get('reference', '')

        prompt = f"""
        Com base EXCLUSIVAMENTE nas notas de pesquisa fornecidas e no texto de {reference}, crie um esboço de sermão expositivo completo.

        O sermão deve ser Cristocêntrico e seguir a estrutura homilética padrão.

        Retorne APENAS um JSON com esta estrutura exata:
        {{
            "ict": "Ideia Central do Texto (uma frase)",
            "tese": "A verdade atemporal para hoje",
            "proposito_basico": "Doutrinário, Evangelístico, Consolação, Ético ou Consagração",
            "proposito_especifico": "O que a igreja deve fazer/sentir?",
            "intro": "Texto da introdução (gancho + transição)",
            "topicos": [
                {{
                    "titulo": "Título do Tópico 1",
                    "explicacao": "Exegese e explicação teológica",
                    "ilustracao": "Sugestão de ilustração breve",
                    "aplicacao": "Aplicação prática para a igreja"
                }},
                {{
                    "titulo": "Título do Tópico 2",
                    "explicacao": "...",
                    "ilustracao": "...",
                    "aplicacao": "..."
                }}
            ],
            "conclusao": "Texto da conclusão e apelo"
        }}
        """

        # Usa o JSON Mode do Groq e o System Prompt de Homilética
        json_str = generate_with_groq(prompt, notes, system_prompt=HOMILETICS_SYSTEM_PROMPT, json_mode=True)

        # Parsing seguro
        try:
            sermon_data = json.loads(json_str)
        except:
            # Fallback se vier markdown ```json
            clean = re.sub(r'```json\s*|\s*```', '', json_str)
            sermon_data = json.loads(clean)

        return jsonify(sermon_data)

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# --- 7. ROTAS DE DOWNLOAD ---

@app.route('/download_notes', methods=['POST'])
def download_notes():
    try:
        content = request.form.get('notes_content', '')
        document = Document()
        document.add_heading('Exposibot - Notas de Pesquisa (Referenciadas)', 0)

        clean_text = re.sub('<[^<]+?>', '', content)
        document.add_paragraph(clean_text)

        f = BytesIO()
        document.save(f)
        f.seek(0)
        return send_file(f, as_attachment=True, download_name='notas_exposibot.docx')
    except Exception as e:
        return f"Erro ao gerar arquivo: {str(e)}", 500

@app.route('/download_sermon_docx', methods=['POST'])
def download_sermon_docx():
    try:
        data = request.form
        document = Document()
        style = document.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)

        document.add_heading('Esboço de Pregação Expositiva', 0)
        document.add_paragraph(f"Texto Base: {data.get('texto_biblico')}")
        document.add_paragraph(f"ICT: {data.get('ict')}")
        document.add_paragraph(f"Tese: {data.get('tese')}")
        document.add_paragraph(f"Propósito: {data.get('proposito_basico')} | {data.get('proposito_especifico')}")
        document.add_paragraph("-" * 60)

        document.add_heading('Introdução', level=1)
        document.add_paragraph(data.get('intro'))

        titulos = data.getlist('topico_titulo[]')
        explicacoes = data.getlist('topico_explicacao[]')
        ilustracoes = data.getlist('topico_ilustracao[]')
        aplicacoes = data.getlist('topico_aplicacao[]')

        for i in range(len(titulos)):
            document.add_heading(f"{i+1}. {titulos[i]}", level=2)
            document.add_paragraph("Explicação:", style='Strong')
            document.add_paragraph(explicacoes[i])
            if ilustracoes[i]:
                document.add_paragraph("Ilustração:", style='Strong')
                try:
                    p = document.add_paragraph(ilustracoes[i])
                    p.style = document.styles['Quote']
                except:
                    document.add_paragraph(ilustracoes[i])
            document.add_paragraph("Aplicação:", style='Strong')
            document.add_paragraph(aplicacoes[i])

        document.add_heading('Conclusão', level=1)
        document.add_paragraph(data.get('conclusao'))

        f = BytesIO()
        document.save(f)
        f.seek(0)
        return send_file(f, as_attachment=True, download_name='sermao_exposibot.docx')
    except Exception as e:
        return f"Erro ao gerar sermão: {str(e)}", 500

if __name__ == '__main__':
    app.run(debug=True)